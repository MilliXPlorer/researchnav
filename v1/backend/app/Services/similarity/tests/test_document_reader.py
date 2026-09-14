from pathlib import Path
import warnings
from unittest.mock import MagicMock, patch

import pytest
from docx import Document
from pypdf import PdfWriter

from similarity.document_reader import DocumentReadError, _parts_from_units, extract_text


def test_extract_text_reads_docx_paragraphs_and_tables(tmp_path: Path) -> None:
    path = tmp_path / "private-manuscript.docx"
    document = Document()
    document.add_paragraph("Research findings")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Table evidence"
    document.save(str(path))

    assert extract_text(path) == "Research findings\nTable evidence"


@pytest.mark.parametrize("suffix", [".pdf", ".docx"])
def test_extract_text_returns_an_empty_string_for_valid_empty_documents(
    tmp_path: Path, suffix: str
) -> None:
    path = tmp_path / f"private-empty{suffix}"
    if suffix == ".pdf":
        writer = PdfWriter()
        writer.add_blank_page(width=72, height=72)
        with path.open("wb") as stream:
            writer.write(stream)
    else:
        Document().save(str(path))

    assert extract_text(path) == ""


@pytest.mark.parametrize("suffix", [".pdf", ".docx"])
def test_malformed_documents_do_not_disclose_private_paths(
    tmp_path: Path, suffix: str, caplog: pytest.LogCaptureFixture
) -> None:
    path = tmp_path / f"private-manuscript{suffix}"
    path.write_bytes(b"not a document")

    with pytest.raises(DocumentReadError) as raised:
        extract_text(path)

    assert str(raised.value) == "Document could not be read."
    assert path.name not in str(raised.value)
    assert not caplog.records


def test_pdf_parser_warnings_are_suppressed_at_the_reader_boundary() -> None:
    document = MagicMock()
    document.__enter__.return_value = document
    document.needs_pass = False
    document.page_count = 0

    def warned_open(**_kwargs: object) -> MagicMock:
        warnings.warn("malformed private PDF", UserWarning)
        return document

    with warnings.catch_warnings(record=True) as caught, patch(
        "similarity.document_reader.pymupdf.open", side_effect=warned_open
    ):
        assert extract_text_for_pdf_reader() == ""

    assert caught == []


def test_incremental_parts_preserve_all_units_in_exact_order() -> None:
    units = ["Chapter One\nFirst paragraph", "Second paragraph", "Final words"]

    parts = _parts_from_units(units, 20)

    assert len(parts) > 1
    assert "".join(parts) == "Chapter One First paragraph Second paragraph Final words"
    assert all(part.endswith(" ") for part in parts[:-1])


def extract_text_for_pdf_reader() -> str:
    """Exercise the PDF boundary without creating a real private file."""
    from io import BytesIO
    from similarity.document_reader import _pdf_text

    return _pdf_text(BytesIO(b"malformed"))
